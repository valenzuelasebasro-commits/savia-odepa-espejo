#!/usr/bin/env python3
"""
Espejo ODEPA — descubrimiento, descarga y base histórica.  v2.0 (28-sep-2026)

Reemplaza la capa de ingesta de `ingest.js` (que dependía del navegador:
DOMParser, DecompressionStream, IndexedDB) por una equivalente en Python que
corre en un runner de GitHub Actions.  El motor de análisis (`analyze.js`) NO se
porta: es puro y se usa tal cual desde Node, así que las reglas de serie de
referencia, alertas y puntaje de riesgo quedan byte a byte idénticas.

REGLA NÚMERO UNO (runbook): no adivinar el nombre del archivo.
  1. Descubrir, no construir: se lee la página del boletín y se toma CUALQUIER
     enlace .xlsx, sin importar cómo se llame.
  2. Fechar por 8 dígitos, no por plantilla (`fecha_desde_nombre`).
  3. Validar por contenido, no por metadatos: todo .xlsx es un ZIP -> 'PK'.
     ODEPA sirve su página de error con status 200, el content-type miente.
  4. Fallar fuerte: si la página no entrega enlaces, salir con código != 0.
  5. Los patrones fijos son SOLO último recurso, y solo para días pasados que la
     página ya no enlaza (subcomando `backfill`).  El camino diario nunca los usa.

Subcomandos
  discover              Lista lo que la página enlaza hoy (no descarga).  Falla si 0.
  fetch                 Descubre y descarga lo nuevo a data/.  Falla si 0 enlaces.
  backfill --from --to  ÚLTIMO RECURSO: sondea patrones conocidos para días pasados.
  parse                 Parsea data/*.xlsx hacia el store (idempotente).
  status                Resumen del store.
"""
from __future__ import annotations

import argparse
import datetime as dt
import io
import json
import os
import re
import sys
import unicodedata
import urllib.error
import urllib.request
import zipfile

HOST = "https://www.odepa.gob.cl"
PAGINA_BOLETIN = (
    "/publicaciones/boletines/"
    "boletin-diario-de-precios-y-volumenes-de-frutas-en-mercados-mayoristas"
)
UA = "Mozilla/5.0 (compatible; SaviaOdepaMirror/2.0; +https://github.com)"

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DATA = os.path.join(ROOT, "data")
STORE = os.path.join(ROOT, "store")
OBS_PATH = os.path.join(STORE, "obs.ndjson")
LOADS_PATH = os.path.join(STORE, "loads.json")
RUNS_PATH = os.path.join(STORE, "runs.json")
URLS_PATH = os.path.join(STORE, "urls.json")  # nombre de archivo -> ruta en odepa.gob.cl

# ---------------------------------------------------------------------------
# Normalización — port literal de ingest.js
# ---------------------------------------------------------------------------
MARKETS = {
    "lo valledor": "Lo Valledor",
    "vega central mapocho": "Vega Central Mapocho",
    "mapocho vta.dir": "Mapocho Venta Directa",
    "macroferia talca": "Macroferia Talca",
    "femacal": "Femacal",
    "la palmera": "La Palmera",
    "solcoagro": "Solcoagro",
    "vega monumental": "Vega Monumental",
    "lagunita pto.montt": "Lagunitas Pto. Montt",
    "vega modelo temuco": "Vega Modelo Temuco",
    "agrochillan": "Agro Chillán",
    "agronor": "Agronor",
}
PRODUCT_ALIAS = {
    "arándano (blue)": "Arándano",
    "arandano": "Arándano",
    "palta": "Palta",
    "zapallo guarda": "Zapallo camote",
}

_WS = re.compile(r"\s+")


def clean(s) -> str:
    if s is None:
        return ""
    return _WS.sub(" ", str(s)).strip()


def cap(s: str) -> str:
    """Equivalente a s.charAt(0).toUpperCase() + s.slice(1): solo el 1.er carácter."""
    return (s[0].upper() + s[1:]) if s else s


def norm_p(s) -> str:
    c = clean(s)
    return PRODUCT_ALIAS.get(c.lower(), cap(c))


_VAR_SUF = re.compile(r"\s*\((?:o|a|os|as)\)\s*$", re.I)


def norm_v(s) -> str:
    return cap(_VAR_SUF.sub("", clean(s)))


_UNIT_DOLLAR = re.compile(r"\s*\$\s*/\s*")


def norm_u(s) -> str:
    return _UNIT_DOLLAR.sub("$/", clean(s).lower(), count=1)


def num(x):
    """Port de `num`: número tal cual; texto con '.' de miles y ',' decimal."""
    if isinstance(x, bool):
        return None
    if isinstance(x, (int, float)):
        return float(x)
    if x is None or x == "":
        return None
    try:
        return float(str(x).replace(".", "").replace(",", "."))
    except ValueError:
        return None


def fold(s: str) -> str:
    return "".join(
        ch for ch in unicodedata.normalize("NFD", str(s)) if unicodedata.category(ch) != "Mn"
    ).lower().strip()


# ---------------------------------------------------------------------------
# Fecha desde el nombre — port literal de fechaDesdeNombre()
# ---------------------------------------------------------------------------
_EIGHT = re.compile(r"(?<!\d)\d{8}(?!\d)")


def fecha_desde_nombre(nombre: str):
    """Devuelve 'AAAAMMDD' canónico, o None. Tolera _20260923, -23092026, 23092026…"""
    s = nombre
    try:
        s = urllib.parse.unquote(nombre)
    except Exception:
        pass
    for b in _EIGHT.findall(s):
        y1, y2 = int(b[:4]), int(b[4:])
        if 2015 <= y1 <= 2035:  # AAAAMMDD
            mo, da = int(b[4:6]), int(b[6:8])
            if 1 <= mo <= 12 and 1 <= da <= 31:
                return b
        if 2015 <= y2 <= 2035:  # DDMMAAAA
            da, mo = int(b[:2]), int(b[2:4])
            if 1 <= mo <= 12 and 1 <= da <= 31:
                return f"{y2}{mo:02d}{da:02d}"
    return None


import urllib.parse  # noqa: E402  (usado por fecha_desde_nombre)


# ---------------------------------------------------------------------------
# Red
# ---------------------------------------------------------------------------
def _get(url: str, timeout: int = 60) -> bytes:
    req = urllib.request.Request(url, headers={"User-Agent": UA, "Accept": "*/*"})
    with urllib.request.urlopen(req, timeout=timeout) as r:
        return r.read()


def indexar_pagina() -> dict:
    """{'AAAAMMDD': url} a partir de los href .xlsx de la página. Descubrir, no construir."""
    html = _get(HOST + PAGINA_BOLETIN).decode("utf-8", "replace")
    idx: dict[str, str] = {}
    for m in re.finditer(r"""href=["']([^"']+?\.xlsx)["']""", html, re.I):
        u = m.group(1)
        if re.match(r"^https?:", u, re.I):
            try:
                u = urllib.parse.urlsplit(u).path
            except Exception:
                pass
        d = fecha_desde_nombre(u.split("/")[-1])
        if d:
            idx[d] = u
    return idx


def bajar_si_xlsx(url: str):
    """Descarga y valida que sea un ZIP real ('PK'). No confía en el content-type."""
    if url.startswith("/"):
        url = HOST + url
    try:
        buf = _get(url)
    except (urllib.error.HTTPError, urllib.error.URLError, TimeoutError):
        return None
    if len(buf) < 4 or buf[0] != 0x50 or buf[1] != 0x4B:
        return None
    return buf


# Último recurso (runbook regla 1.5): SOLO para días pasados que la página ya no
# enlaza. Si ODEPA renombra otra vez, esto se rompe — el descubrimiento no.
PATRONES = [
    lambda d: f"BoletinDiarioFrutasHortalizas-{d[6:8]}{d[4:6]}{d[:4]}.xlsx",
    lambda d: f"BoletinDiarioFrutasHortalizas_{d[6:8]}{d[4:6]}{d[:4]}.xlsx",
    lambda d: f"BoletinDiarioFrutasHortalizas_{d}.xlsx",
    lambda d: f"Boletin_Diario_de_Frutas_y_Hortalizas_{d}.xlsx",
]


_ULTIMO_PATRON = [0]  # índice del último patrón que funcionó (días seguidos suelen compartirlo)


def sondear(d: str):
    """d = AAAAMMDD. Prueba patrones conocidos en los uploads del mes ±1.

    Empieza por el patrón que funcionó en el día anterior: ODEPA alterna nombres,
    pero en rachas, así que esto evita ~9 intentos fallidos por día en un backfill
    largo (que si no, excede el límite de tiempo del job)."""
    y, m = int(d[:4]), int(d[4:6])
    meses = [
        (y, m),
        (y - 1 if m == 1 else y, 12 if m == 1 else m - 1),
        (y + 1 if m == 12 else y, 1 if m == 12 else m + 1),
    ]
    primero = _ULTIMO_PATRON[0]
    orden = [primero] + [i for i in range(len(PATRONES)) if i != primero]
    for i in orden:
        fname = PATRONES[i](d)
        for fy, fm in meses:
            url = f"/wp-content/uploads/{fy}/{fm:02d}/{fname}"
            buf = bajar_si_xlsx(url)
            if buf:
                _ULTIMO_PATRON[0] = i
                return url, fname, buf
    return None


# ---------------------------------------------------------------------------
# Parseo del boletín — port de parseBoletin() usando openpyxl
# ---------------------------------------------------------------------------
_SHEET_RE = re.compile(r"^(Frutas|Hortalizas)_(.+)$")
_DIA_RE = re.compile(r"^d[ií]a:?$", re.I)
_PROD_RE = re.compile(r"^producto$", re.I)

EXCEL_EPOCH = dt.date(1899, 12, 30)


def _xl_date(v):
    """Serial de Excel o datetime -> 'AAAA-MM-DD'."""
    if isinstance(v, dt.datetime):
        return v.date().isoformat()
    if isinstance(v, dt.date):
        return v.isoformat()
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        return (EXCEL_EPOCH + dt.timedelta(days=round(v))).isoformat()
    return None


def _headers(row) -> dict:
    hdr: dict[str, int] = {}
    for i, h in enumerate(row):
        k = fold(clean(h))
        if not k:
            continue
        if k.startswith("producto"):
            hdr["p"] = i
        elif k.startswith("variedad"):
            hdr["v"] = i
        elif k.startswith("calidad"):
            hdr["c"] = i
        elif k.startswith("volumen"):
            hdr["vol"] = i
        elif "maximo" in k:
            hdr["max"] = i
        elif "minimo" in k:
            hdr["min"] = i
        elif "promedio" in k:
            hdr["avg"] = i
        elif k.startswith("unidad"):
            hdr["u"] = i
        elif k.startswith("origen"):
            hdr["o"] = i
    return hdr


def parse_boletin(buf: bytes, file: str):
    """Devuelve (data_date, rows) con rows ya consolidadas por llave."""
    import openpyxl

    wb = openpyxl.load_workbook(io.BytesIO(buf), read_only=True, data_only=True)
    rows = []
    data_date = None
    try:
        for name in wb.sheetnames:
            mm = _SHEET_RE.match(name)
            if not mm:
                continue
            sub = "Frutas" if mm.group(1) == "Frutas" else "Hortalizas y tubérculos"
            raw_market = mm.group(2).strip()
            market = MARKETS.get(raw_market.lower(), clean(raw_market))
            ws = wb[name]
            hdr, fecha = None, None
            for r in ws.iter_rows(values_only=True):
                if not r:
                    continue
                c0 = clean(r[0] if len(r) else None)
                if _DIA_RE.match(c0) and len(r) > 1:
                    d = _xl_date(r[1])
                    if d:
                        fecha = d
                    continue
                if _PROD_RE.match(c0):
                    hdr = _headers(r)
                    continue
                if not hdr or not fecha or not c0:
                    continue
                if "avg" not in hdr or hdr["avg"] >= len(r):
                    continue
                pavg = num(r[hdr["avg"]])
                if pavg is None or not pavg > 0:
                    continue

                def cell(key):
                    i = hdr.get(key)
                    return r[i] if i is not None and i < len(r) else None

                pmin = num(cell("min"))
                pmax = num(cell("max"))
                rows.append(
                    {
                        "f": fecha,
                        "m": market,
                        "sub": sub,
                        "p": norm_p(cell("p")),
                        "v": norm_v(cell("v")),
                        "c": cap(clean(cell("c"))),
                        "u": norm_u(cell("u")),
                        "o": clean(cell("o")),
                        "vol": num(cell("vol")) or 0.0,
                        "pmin": pavg if pmin is None else pmin,
                        "pmax": pavg if pmax is None else pmax,
                        "pavg": pavg,
                    }
                )
            if fecha:
                data_date = fecha if (data_date is None or fecha > data_date) else data_date
    finally:
        wb.close()

    # Consolidar duplicados dentro del mismo archivo (misma llave)
    out: dict[str, dict] = {}
    for r in rows:
        k = "|".join([r["f"], r["m"], r["p"], r["v"], r["c"], r["u"], r["o"]])
        e = out.get(k)
        if e is None:
            out[k] = {"k": k, **r, "src": file}
        else:
            w1 = e["vol"] or 1
            w2 = r["vol"] or 1
            e["pavg"] = (e["pavg"] * w1 + r["pavg"] * w2) / (w1 + w2)
            e["vol"] += r["vol"]
            e["pmin"] = min(e["pmin"], r["pmin"])
            e["pmax"] = max(e["pmax"], r["pmax"])
    return data_date, list(out.values())


# ---------------------------------------------------------------------------
# Store versionado en el repo (reemplaza IndexedDB)
# ---------------------------------------------------------------------------
def _load_json(path, default):
    if not os.path.exists(path):
        return default
    with open(path, encoding="utf-8") as fh:
        return json.load(fh)


def _save_json(path, obj):
    os.makedirs(os.path.dirname(path), exist_ok=True)
    tmp = path + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(obj, fh, ensure_ascii=False, indent=1, sort_keys=True)
        fh.write("\n")
    os.replace(tmp, path)


def load_obs() -> dict:
    obs: dict[str, dict] = {}
    if os.path.exists(OBS_PATH):
        with open(OBS_PATH, encoding="utf-8") as fh:
            for line in fh:
                line = line.strip()
                if line:
                    o = json.loads(line)
                    obs[o["k"]] = o
    return obs


def save_obs(obs: dict):
    os.makedirs(STORE, exist_ok=True)
    tmp = OBS_PATH + ".tmp"
    with open(tmp, "w", encoding="utf-8") as fh:
        for k in sorted(obs):  # orden estable => diffs legibles en git
            fh.write(json.dumps(obs[k], ensure_ascii=False, sort_keys=True) + "\n")
    os.replace(tmp, OBS_PATH)


def upsert(obs: dict, rows: list, ts: str):
    """Port de upsert(): rev/prev/ts igual que en IndexedDB."""
    ins = upd = same = 0
    for r in rows:
        old = obs.get(r["k"])
        if old is None:
            obs[r["k"]] = {**r, "rev": 1, "ts": ts}
            ins += 1
        elif (
            old.get("vol") == r["vol"]
            and old.get("pavg") == r["pavg"]
            and old.get("pmin") == r["pmin"]
            and old.get("pmax") == r["pmax"]
        ):
            same += 1
        else:
            obs[r["k"]] = {
                **r,
                "rev": (old.get("rev") or 1) + 1,
                "ts": ts,
                "prev": {
                    "src": old.get("src"),
                    "vol": old.get("vol"),
                    "pavg": old.get("pavg"),
                    "ts": old.get("ts"),
                },
            }
            upd += 1
    return ins, upd, same


# ---------------------------------------------------------------------------
# Subcomandos
# ---------------------------------------------------------------------------
def _fail(msg: str, code: int = 2):
    print(f"::error::{msg}", file=sys.stderr)
    print(f"FALLA FUERTE: {msg}", file=sys.stderr)
    sys.exit(code)


def cmd_discover(_args):
    try:
        idx = indexar_pagina()
    except Exception as e:
        _fail(f"no se pudo leer la página del boletín: {e!r}")
        return
    if not idx:
        _fail(
            "la página del boletín no entregó ningún enlace .xlsx. "
            "ODEPA cambió la estructura del sitio: revisar el selector antes de confiar en el espejo."
        )
    for d in sorted(idx):
        print(f"{d}\t{idx[d]}")
    print(f"\nenlaces descubiertos: {len(idx)}")


def cmd_fetch(args):
    try:
        idx = indexar_pagina()
    except Exception as e:
        _fail(f"no se pudo leer la página del boletín: {e!r}")
        return
    if not idx:
        _fail(
            "la página del boletín no entregó ningún enlace .xlsx. "
            "ODEPA cambió la estructura del sitio: el espejo NO se degrada en silencio."
        )
    os.makedirs(DATA, exist_ok=True)
    urls = _load_json(URLS_PATH, {})
    nuevos, ya, fallidos = [], [], []
    for d in sorted(idx):
        url = idx[d]
        fname = url.split("/")[-1]
        dest = os.path.join(DATA, fname)
        urls[fname] = {"url": url, "via": "pagina"}
        if os.path.exists(dest) and not args.force:
            ya.append(fname)
            continue
        buf = bajar_si_xlsx(url)
        if buf is None:
            fallidos.append((d, url))
            continue
        with open(dest, "wb") as fh:
            fh.write(buf)
        nuevos.append((d, fname, len(buf)))
        print(f"descargado {fname} ({len(buf)} bytes) fecha={d} via=pagina")
    _save_json(URLS_PATH, urls)
    for f in ya:
        print(f"ya estaba: {f}")
    if fallidos:
        for d, u in fallidos:
            print(f"::warning::enlace no descargable o no-ZIP: {d} {u}", file=sys.stderr)
    print(f"\nindicePagina={len(idx)} nuevos={len(nuevos)} yaEstaban={len(ya)} fallidos={len(fallidos)}")
    if not nuevos and not ya:
        _fail("había enlaces pero no se pudo descargar ninguno como XLSX válido")


def cmd_backfill(args):
    """ÚLTIMO RECURSO. La página no tiene archivo histórico."""
    os.makedirs(DATA, exist_ok=True)
    urls = _load_json(URLS_PATH, {})
    a = dt.date.fromisoformat(args.from_)
    b = dt.date.fromisoformat(args.to)
    hit = miss = 0
    d = a
    while d <= b:
        if d.weekday() == 6:  # domingo
            d += dt.timedelta(days=1)
            continue
        key = d.strftime("%Y%m%d")
        found = sondear(key)
        if found:
            url, fname, buf = found
            urls[fname] = {"url": url, "via": "sondeo"}
            dest = os.path.join(DATA, fname)
            if not os.path.exists(dest) or args.force:
                with open(dest, "wb") as fh:
                    fh.write(buf)
                print(f"sondeo OK {d} -> {fname} ({len(buf)} bytes)", flush=True)
            hit += 1
        else:
            print(f"sin boletín {d}", flush=True)
            miss += 1
        d += dt.timedelta(days=1)
    _save_json(URLS_PATH, urls)
    print(f"\nbackfill {args.from_}..{args.to}: encontrados={hit} sin_boletin={miss}")
    print("NOTA: el sondeo por patrones es la parte frágil. El camino diario usa descubrimiento.")


def cmd_parse(args):
    obs = load_obs()
    loads = _load_json(LOADS_PATH, {})
    runs = _load_json(RUNS_PATH, [])
    urls = _load_json(URLS_PATH, {})
    ts = dt.datetime.now(dt.timezone.utc).isoformat()
    files = sorted(f for f in os.listdir(DATA) if f.lower().endswith(".xlsx")) if os.path.isdir(DATA) else []
    if not files:
        _fail("no hay archivos .xlsx en data/: nada que parsear")
    tot = {"ins": 0, "upd": 0, "same": 0, "ok": 0, "err": 0, "skip": 0, "skip_err": 0}
    for fname in files:
        prev = loads.get(fname)
        path = os.path.join(DATA, fname)
        size = os.path.getsize(path)
        if prev and prev.get("bytes") == size and not args.force:
            # Ya visto con el mismo tamaño: no se reintenta.
            if prev.get("status") == "ok":
                tot["skip"] += 1
                continue
            if prev.get("status") == "error":
                # Ya se falló en rojo por este archivo en su momento; no se repite
                # el rojo cada día, pero tampoco se oculta.
                tot["skip_err"] += 1
                print(f"::warning::{fname} sigue ilegible (error ya registrado): {prev.get('error')}",
                      file=sys.stderr)
                continue
        try:
            with open(path, "rb") as fh:
                buf = fh.read()
            data_date, rows = parse_boletin(buf, fname)
            if not rows:
                raise ValueError("0 filas útiles: ¿cambió el layout de las hojas?")
            ins, upd, same = upsert(obs, rows, ts)
            pub = fecha_desde_nombre(fname)
            meta = urls.get(fname) or {}
            loads[fname] = {
                "file": fname,
                "url": meta.get("url"),
                "via": meta.get("via") or (prev.get("via") if prev else "pagina"),
                "pub": f"{pub[:4]}-{pub[4:6]}-{pub[6:8]}" if pub else data_date,
                "dataDate": data_date,
                "rows": len(rows),
                "bytes": size,
                "ins": ins,
                "upd": upd,
                "same": same,
                "status": "ok",
                "ts": ts,
                "loads": (prev.get("loads", 1) if prev else 0) + 1,
            }
            tot["ins"] += ins
            tot["upd"] += upd
            tot["same"] += same
            tot["ok"] += 1
            print(f"{fname}: fecha={data_date} filas={len(rows)} nuevas={ins} act={upd} igual={same}")
        except Exception as e:
            loads[fname] = {"file": fname, "status": "error", "error": repr(e), "ts": ts, "bytes": size}
            tot["err"] += 1
            print(f"::error::{fname}: {e!r}", file=sys.stderr)
    save_obs(obs)
    _save_json(LOADS_PATH, loads)
    runs.append({"ts": ts, "action": "parse", **tot, "registros": len(obs)})
    _save_json(RUNS_PATH, runs[-200:])
    print(
        f"\nstore: registros={len(obs)} archivos_ok={sum(1 for l in loads.values() if l.get('status') == 'ok')} "
        f"nuevas={tot['ins']} act={tot['upd']} igual={tot['same']} omitidos={tot['skip']} "
        f"ilegibles_ya_conocidos={tot['skip_err']} errores_nuevos={tot['err']}"
    )
    # El store ya quedó guardado arriba: los datos buenos nunca se pierden por un
    # archivo malo. Pero un archivo NUEVO que no se pudo parsear deja la corrida en
    # rojo, porque suele significar que ODEPA cambió el layout de las hojas.
    if tot["err"]:
        _fail(
            f"{tot['err']} archivo(s) nuevo(s) no se pudieron parsear. "
            "Si es el boletín del día, probablemente cambió el layout de las hojas "
            "(nombres 'Frutas_<mercado>' / 'Hortalizas_<mercado>', fila 'Día:', cabecera 'Producto'). "
            "El store quedó intacto con los datos anteriores."
        )


def cmd_status(_args):
    obs = load_obs()
    loads = _load_json(LOADS_PATH, {})
    ok = sorted(
        (l for l in loads.values() if l.get("status") == "ok"),
        key=lambda l: l.get("dataDate") or "",
    )
    print(
        json.dumps(
            {
                "registros": len(obs),
                "archivos": len(ok),
                "primera": ok[0]["dataDate"] if ok else None,
                "ultima": ok[-1]["dataDate"] if ok else None,
                "errores": sum(1 for l in loads.values() if l.get("status") == "error"),
            },
            ensure_ascii=False,
            indent=2,
        )
    )


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    sub = ap.add_subparsers(dest="cmd", required=True)
    sub.add_parser("discover").set_defaults(fn=cmd_discover)
    f = sub.add_parser("fetch")
    f.add_argument("--force", action="store_true")
    f.set_defaults(fn=cmd_fetch)
    b = sub.add_parser("backfill")
    b.add_argument("--from", dest="from_", required=True)
    b.add_argument("--to", required=True)
    b.add_argument("--force", action="store_true")
    b.set_defaults(fn=cmd_backfill)
    p = sub.add_parser("parse")
    p.add_argument("--force", action="store_true")
    p.set_defaults(fn=cmd_parse)
    sub.add_parser("status").set_defaults(fn=cmd_status)
    args = ap.parse_args()
    args.fn(args)


if __name__ == "__main__":
    main()
